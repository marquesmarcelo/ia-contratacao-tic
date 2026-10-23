#!/usr/bin/env python3
"""Despeja uma planilha .xlsx em markdown legível: por aba, cada linha com as
células não vazias no formato `A12="rótulo"`, `C12=123` ou `D12=<fórmula: =B12*C12>`.

Serve para os agentes entenderem o template de TCO (onde ficam os rótulos,
quais células são de entrada e quais são calculadas) sem abrir o Excel.

Uso: python3 despejar_xlsx.py <arquivo.xlsx> [saida.md] [--max-linhas N]

Abas com mais de N linhas (padrão 2000) são cortadas e o corte é anotado.
"""
import sys

import openpyxl


def main():
    args = sys.argv[1:]
    limite = 2000
    if "--max-linhas" in args:
        i = args.index("--max-linhas")
        limite = int(args[i + 1])
        del args[i:i + 2]
    origem = args[0]
    wb = openpyxl.load_workbook(origem)  # fórmulas preservadas (data_only=False)
    linhas = [f"# {origem}", "",
              "Células com `<fórmula: ...>` são calculadas: não sobrescrever. "
              "Demais células numéricas ou vazias ao lado de rótulos são de entrada.", ""]
    for ws in wb.worksheets:
        linhas += [f"## Aba: {ws.title}", ""]
        for n, row in enumerate(ws.iter_rows(), start=1):
            if n > limite:
                linhas.append(f"- [aba cortada: {ws.max_row} linhas no total, mostradas as primeiras {limite}; consulte o arquivo original]")
                break
            partes = []
            for c in row:
                v = c.value
                if v is None:
                    continue
                if isinstance(v, str) and v.startswith("="):
                    partes.append(f"{c.coordinate}=<fórmula: {v}>")
                elif isinstance(v, str):
                    partes.append(f'{c.coordinate}="{v.strip()}"')
                else:
                    partes.append(f"{c.coordinate}={v}")
            if partes:
                linhas.append("- " + " | ".join(partes))
        linhas.append("")
    texto = "\n".join(linhas)
    if len(args) > 1:
        open(args[1], "w", encoding="utf-8").write(texto)
    else:
        print(texto)


if __name__ == "__main__":
    main()
