#!/usr/bin/env python3
"""Gera o template do Mapa de Preços (IN SEGES/ME 65/2021; IN SGD/ME 94/2022, art. 20).

Premissas:
- Até 30 itens no Mapa e 300 linhas de coleta; ampliar copiando linhas.
- Média, mediana, menor valor, desvio padrão e CV consideram só preços com situação "Válido".
- Mediana e desvio usam fórmula matricial (MEDIAN(IF(...))), aceita por Excel e LibreOffice.
Uso: python3 -I gerar_mapa_precos.py <saida.xlsx>
"""
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula

FONTE = "Arial"
ENTRADA = PatternFill("solid", fgColor="FFF4CC")   # células para preencher
CALC = PatternFill("solid", fgColor="E7EEF5")      # células calculadas
CAB = PatternFill("solid", fgColor="1F4E5A")
EXEMPLO = Font(name=FONTE, size=10, italic=True, color="7F7F7F")
FINA = Side(style="thin", color="BFBFBF")
BORDA = Border(left=FINA, right=FINA, top=FINA, bottom=FINA)
N_ITENS, N_COLETA, N_FORN = 30, 300, 40
MOEDA = 'R$ #,##0.00'


def cabecalho(ws, linha, titulos, larguras):
    for i, (t, w) in enumerate(zip(titulos, larguras), start=1):
        c = ws.cell(row=linha, column=i, value=t)
        c.font = Font(name=FONTE, size=10, bold=True, color="FFFFFF")
        c.fill = CAB
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BORDA
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[linha].height = 42
    ws.freeze_panes = ws.cell(row=linha + 1, column=1)


def orientacoes(wb):
    ws = wb.active
    ws.title = "Orientações"
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 110
    linhas = [
        ("Mapa de Preços — template", True),
        ("Documento da pesquisa de preços da contratação (IN SEGES/ME nº 65/2021, art. 3º; IN SGD/ME nº 94/2022, art. 20). "
         "Responsáveis: Integrante Administrativo (coleta, fontes, fornecedores) com o Integrante Técnico (análise crítica e preço estimado).", False),
        ("", False),
        ("Legenda", True),
        ("Fundo amarelo: célula para preencher.   Fundo azul: célula calculada, não editar.   Texto cinza em itálico: linha de exemplo, apagar antes de usar.", False),
        ("", False),
        ("Como preencher", True),
        ("1. Aba Identificação: objeto, processo, equipe e datas (art. 3º, I e II).", False),
        ("2. Aba Itens: um item por linha, com unidade, quantidade (do ETP) e o método estatístico escolhido.", False),
        ("3. Aba Coleta: um preço por linha, com o parâmetro do art. 5º (I a V), a fonte, a data e o documento de suporte. "
         "Marque como Excluído o preço inexequível, inconsistente ou excessivamente elevado e escreva o motivo (art. 6º, caput e § 3º).", False),
        ("4. Aba Fornecedores: todos os fornecedores consultados na pesquisa direta, inclusive os que não responderam (art. 5º, § 2º, IV).", False),
        ("5. Aba Mapa: confira as estatísticas e os alertas. O preço estimado sai do método escolhido na aba Itens, com o ajuste percentual do art. 6º, § 2º, se houver.", False),
        ("6. Aba Justificativas: metodologia, critérios de exclusão, escolha dos fornecedores e exceções (art. 3º, VI e VIII; art. 6º, §§ 1º e 5º).", False),
        ("", False),
        ("Regras da IN SEGES/ME 65/2021 refletidas na planilha (conferir no texto da norma antes de usar)", True),
        ("Art. 5º: parâmetros I (sistemas oficiais, ex.: Painel de Preços), II (contratações similares da Administração, até 1 ano), "
         "III (mídia especializada, tabelas e sítios, até 6 meses), IV (pesquisa direta com no mínimo 3 fornecedores, até 6 meses), V (notas fiscais eletrônicas, até 1 ano).", False),
        ("Art. 5º, § 1º: priorizar os parâmetros I e II; se não forem usados, justificar. A coluna de alerta do Mapa aponta o item sem preço válido desses parâmetros.", False),
        ("Art. 6º: método média, mediana ou menor valor, sobre três ou mais preços válidos. Menos de três exige justificativa e aprovação (§ 5º). Outro método exige justificativa e aprovação (§ 1º).", False),
        ("Art. 6º, § 4º: análise crítica quando houver grande variação. O coeficiente de variação (CV) é calculado por item; o limite de alerta é parâmetro do órgão, na aba Identificação.", False),
        ("Art. 6º, § 6º: se o preço vier só do parâmetro I, não pode passar da mediana do item nos sistemas consultados.", False),
        ("IN 94, art. 20, §§ 1º a 3º: proposta de fornecedor só com justificativa; considerar preços praticados pelos fabricantes; usar o PMC-TIC quando o item constar em catálogo, salvo pesquisa com valor inferior.", False),
    ]
    for i, (texto, negrito) in enumerate(linhas, start=2):
        c = ws.cell(row=i, column=2, value=texto)
        c.font = Font(name=FONTE, size=12 if i == 2 else 10, bold=negrito)
        c.alignment = Alignment(wrap_text=True, vertical="top")


def identificacao(wb):
    ws = wb.create_sheet("Identificação")
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 70
    campos = [
        ("Órgão/entidade", ""),
        ("Processo nº", ""),
        ("Objeto da contratação (art. 3º, I)", ""),
        ("Responsáveis pela pesquisa / Equipe de Planejamento (art. 3º, II)", ""),
        ("Data de conclusão da pesquisa", ""),
        ("Data prevista de divulgação do edital", ""),
        ("Item em Catálogo de Soluções de TIC com PMC-TIC? (IN 94, art. 20, § 3º)", "Não"),
        ("Limite de coeficiente de variação para alerta (parâmetro do órgão)", 0.25),
    ]
    for i, (rotulo, valor) in enumerate(campos, start=2):
        ws.cell(row=i, column=1, value=rotulo).font = Font(name=FONTE, size=10, bold=True)
        c = ws.cell(row=i, column=2, value=valor)
        c.fill = ENTRADA
        c.border = BORDA
        c.font = Font(name=FONTE, size=10)
    ws["B9"].number_format = "0%"
    ws["B9"].comment = Comment("Valor sugerido apenas como ponto de partida. Defina o limite conforme a orientação do órgão; não é número fixado na IN 65/2021.", "Template")
    dv = DataValidation(type="list", formula1='"Sim,Não"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("B8")
    ws.cell(row=1, column=1, value="Identificação da pesquisa de preços").font = Font(name=FONTE, size=12, bold=True)


def itens(wb):
    ws = wb.create_sheet("Itens")
    cabecalho(ws, 1, ["Item", "Descrição (igual ao TR)", "CATMAT/CATSER", "Unidade", "Quantidade (ETP)",
                      "Método (art. 6º)", "Ajuste % (art. 6º, § 2º)", "PMC-TIC (se houver)"],
              [7, 50, 15, 11, 14, 16, 14, 16])
    dv = DataValidation(type="list", formula1='"Média,Mediana,Menor valor"', allow_blank=True)
    ws.add_data_validation(dv)
    for r in range(2, N_ITENS + 2):
        for col in range(1, 9):
            c = ws.cell(row=r, column=col)
            c.fill = ENTRADA
            c.border = BORDA
        ws.cell(row=r, column=7).number_format = "0.00%"
        ws.cell(row=r, column=8).number_format = MOEDA
        dv.add(f"F{r}")
    exemplo = [1, "EXEMPLO — Notebook corporativo, conforme especificação do TR", "000000", "unidade", 150, "Mediana", 0, None]
    for col, v in enumerate(exemplo, start=1):
        c = ws.cell(row=2, column=col, value=v)
        c.font = EXEMPLO


def coleta(wb):
    ws = wb.create_sheet("Coleta")
    cabecalho(ws, 1, ["Item", "Parâmetro (art. 5º)", "Fonte (identificação: sistema, órgão/contrato, sítio, fornecedor, NF)",
                      "Data do preço / acesso", "Valor unitário coletado", "Fator de atualização (1 = sem)",
                      "Valor unitário atualizado", "Situação", "Motivo da exclusão (art. 6º, § 3º)", "Documento de suporte (arquivo/página)"],
              [7, 15, 46, 14, 16, 13, 16, 12, 40, 34])
    dv_par = DataValidation(type="list", formula1='"I,II,III,IV,V"', allow_blank=True)
    dv_sit = DataValidation(type="list", formula1='"Válido,Excluído"', allow_blank=True)
    ws.add_data_validation(dv_par)
    ws.add_data_validation(dv_sit)
    for r in range(2, N_COLETA + 2):
        for col in (1, 2, 3, 4, 5, 6, 8, 9, 10):
            c = ws.cell(row=r, column=col)
            c.fill = ENTRADA
            c.border = BORDA
        g = ws.cell(row=r, column=7, value=f'=IF(E{r}="","",E{r}*IF(F{r}="",1,F{r}))')
        g.fill = CALC
        g.border = BORDA
        ws.cell(row=r, column=4).number_format = "dd/mm/yyyy"
        ws.cell(row=r, column=5).number_format = MOEDA
        ws.cell(row=r, column=6).number_format = "0.0000"
        g.number_format = MOEDA
        dv_par.add(f"B{r}")
        dv_sit.add(f"H{r}")
    exemplos = [
        [1, "I", "EXEMPLO — Painel de Preços, consulta em dd/mm/aaaa", None, 5200, 1, None, "Válido", "", "insumos/precos/painel.pdf p. 2"],
        [1, "II", "EXEMPLO — Contrato nº xx/aaaa do órgão X", None, 5450, 1.02, None, "Válido", "", "insumos/precos/contrato-x.pdf"],
        [1, "IV", "EXEMPLO — Fornecedor A (cotação formal)", None, 9800, 1, None, "Excluído", "EXEMPLO — 80% acima da mediana dos demais preços", "insumos/precos/cotacao-a.pdf"],
    ]
    for i, linha in enumerate(exemplos, start=2):
        for col, v in enumerate(linha, start=1):
            if col == 7:
                continue
            c = ws.cell(row=i, column=col, value=v)
            c.font = EXEMPLO


def fornecedores(wb):
    ws = wb.create_sheet("Fornecedores")
    cabecalho(ws, 1, ["Fornecedor", "CNPJ/CPF", "E-mail / telefone", "Data da solicitação", "Data da resposta",
                      "Respondeu?", "Justificativa da escolha (art. 3º, VIII)", "Documento (solicitação/proposta)"],
              [32, 20, 30, 15, 15, 12, 46, 34])
    dv = DataValidation(type="list", formula1='"Sim,Não"', allow_blank=True)
    ws.add_data_validation(dv)
    for r in range(2, N_FORN + 2):
        for col in range(1, 9):
            c = ws.cell(row=r, column=col)
            c.fill = ENTRADA
            c.border = BORDA
        ws.cell(row=r, column=4).number_format = "dd/mm/yyyy"
        ws.cell(row=r, column=5).number_format = "dd/mm/yyyy"
        dv.add(f"F{r}")
    for col, v in enumerate(["EXEMPLO — Fornecedor A", "00.000.000/0001-00", "comercial@exemplo.com", None, None, "Sim",
                             "EXEMPLO — revenda autorizada com contratos públicos do mesmo objeto", "insumos/precos/cotacao-a.pdf"], start=1):
        ws.cell(row=2, column=col, value=v).font = EXEMPLO


def mapa(wb):
    ws = wb.create_sheet("Mapa")
    cabecalho(ws, 1, ["Item", "Descrição", "Unidade", "Quantidade", "Nº preços válidos", "Média", "Mediana", "Menor valor",
                      "Desvio padrão", "CV", "Método", "Preço unitário (método)", "Ajuste %", "Preço unitário estimado",
                      "PMC-TIC", "Valor total estimado", "Alertas"],
              [7, 40, 10, 11, 11, 15, 15, 15, 14, 9, 13, 17, 10, 17, 15, 18, 60])
    C = "Coleta!$A$2:$A$%d" % (N_COLETA + 1)
    S = "Coleta!$H$2:$H$%d" % (N_COLETA + 1)
    V = "Coleta!$G$2:$G$%d" % (N_COLETA + 1)
    P = "Coleta!$B$2:$B$%d" % (N_COLETA + 1)
    for r in range(2, N_ITENS + 2):
        i = r  # mesma linha da aba Itens
        f = {
            "A": f'=IF(Itens!A{i}="","",Itens!A{i})',
            "B": f'=IF(A{r}="","",Itens!B{i})',
            "C": f'=IF(A{r}="","",Itens!D{i})',
            "D": f'=IF(A{r}="","",Itens!E{i})',
            "E": f'=IF(A{r}="","",COUNTIFS({C},A{r},{S},"Válido"))',
            "F": f'=IF(OR(A{r}="",E{r}=0),"",AVERAGEIFS({V},{C},A{r},{S},"Válido"))',
            "H": f'=IF(OR(A{r}="",E{r}=0),"",_xlfn.MINIFS({V},{C},A{r},{S},"Válido"))',
            "J": f'=IF(OR(A{r}="",E{r}<2),"",I{r}/F{r})',
            "K": f'=IF(A{r}="","",Itens!F{i})',
            "L": f'=IF(OR(A{r}="",E{r}=0,K{r}=""),"",IF(K{r}="Média",F{r},IF(K{r}="Mediana",G{r},H{r})))',
            "M": f'=IF(A{r}="","",IF(Itens!G{i}="",0,Itens!G{i}))',
            "N": f'=IF(L{r}="","",ROUND(L{r}*(1+M{r}),2))',
            "O": f'=IF(OR(A{r}="",Itens!H{i}=""),"",Itens!H{i})',
            "P": f'=IF(OR(N{r}="",D{r}=""),"",IF(O{r}="",N{r},MIN(N{r},O{r}))*D{r})',
            "Q": (f'=IF(A{r}="","",TRIM('
                  f'IF(E{r}<3,"Menos de 3 preços válidos: justificar e aprovar (art. 6º, § 5º). ","")'
                  f'&IF(COUNTIFS({C},A{r},{S},"Válido",{P},"I")+COUNTIFS({C},A{r},{S},"Válido",{P},"II")=0,"Sem preço dos parâmetros I ou II: justificar (art. 5º, § 1º). ","")'
                  f'&IF(AND(J{r}<>"",J{r}>Identificação!$B$9),"CV acima do limite: análise crítica (art. 6º, § 4º). ","")'
                  f'&IF(K{r}="","Escolher o método na aba Itens. ","")'
                  f'&IF(AND(O{r}<>"",N{r}<>"",N{r}>O{r}),"Estimado acima do PMC-TIC: usado o PMC-TIC (IN 94, art. 20, § 3º). ","")))'),
        }
        for col, formula in f.items():
            ws[f"{col}{r}"] = formula
        ws[f"G{r}"] = ArrayFormula(f"G{r}", f'=IF(OR(A{r}="",E{r}=0),"",MEDIAN(IF(({C}=A{r})*({S}="Válido"),{V})))')
        ws[f"I{r}"] = ArrayFormula(f"I{r}", f'=IF(OR(A{r}="",E{r}<2),"",STDEV(IF(({C}=A{r})*({S}="Válido"),{V})))')
        for col in "ABCDEFGHIJKLMNOPQ":
            c = ws[f"{col}{r}"]
            c.fill = CALC
            c.border = BORDA
            c.alignment = Alignment(wrap_text=(col in "BQ"), vertical="top")
        for col in "FGHILNOP":
            ws[f"{col}{r}"].number_format = MOEDA
        ws[f"J{r}"].number_format = "0.0%"
        ws[f"M{r}"].number_format = "0.00%"
    t = N_ITENS + 3
    ws[f"O{t}"] = "TOTAL"
    ws[f"O{t}"].font = Font(name=FONTE, size=10, bold=True)
    ws[f"P{t}"] = f"=SUM(P2:P{N_ITENS + 1})"
    ws[f"P{t}"].number_format = MOEDA
    ws[f"P{t}"].fill = CALC
    ws[f"P{t}"].font = Font(name=FONTE, size=10, bold=True)


def justificativas(wb):
    ws = wb.create_sheet("Justificativas")
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 90
    topicos = [
        "Caracterização das fontes consultadas (art. 3º, III)",
        "Método estatístico aplicado e justificativa (art. 3º, V e VI; art. 6º)",
        "Critérios para desconsiderar valores inexequíveis, inconsistentes ou excessivamente elevados (art. 6º, § 3º)",
        "Análise crítica dos preços com grande variação (art. 6º, § 4º)",
        "Justificativa da ausência dos parâmetros I e II, se for o caso (art. 5º, § 1º)",
        "Justificativa da escolha dos fornecedores na pesquisa direta (art. 3º, VIII)",
        "Uso de propostas de fornecedores como única fonte, se for o caso (IN 94, art. 20, § 1º)",
        "Itens com menos de 3 preços ou outro método: justificativa e aprovação (art. 6º, §§ 1º e 5º)",
        "Orçamentos fora do prazo e índice de atualização usado (art. 5º, § 3º)",
        "Memória de cálculo e documentos de suporte (art. 3º, VII)",
    ]
    ws.cell(row=1, column=1, value="Justificativas da pesquisa de preços").font = Font(name=FONTE, size=12, bold=True)
    for i, t in enumerate(topicos, start=3):
        a = ws.cell(row=i, column=1, value=t)
        a.font = Font(name=FONTE, size=10, bold=True)
        a.alignment = Alignment(wrap_text=True, vertical="top")
        b = ws.cell(row=i, column=2)
        b.fill = ENTRADA
        b.border = BORDA
        b.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 60


def main():
    wb = Workbook()
    orientacoes(wb)
    identificacao(wb)
    itens(wb)
    coleta(wb)
    fornecedores(wb)
    mapa(wb)
    justificativas(wb)
    wb.save(sys.argv[1])
    print("salvo", sys.argv[1])


if __name__ == "__main__":
    main()
